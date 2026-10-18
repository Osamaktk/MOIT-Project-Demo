from datetime import datetime, timedelta, timezone
from pathlib import Path
from uuid import uuid4

from docx import Document
from fastapi import APIRouter, BackgroundTasks, Depends, File, HTTPException, Request, UploadFile
from fastapi.encoders import jsonable_encoder
from fastapi.responses import FileResponse
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.ai.gateway import AIGateway
from app.ai.stt import FasterWhisperProvider, STTProvider
from app.core.config import settings
from app.core.deps import get_current_user, get_db, require_roles
from app.db.models import AIRun, Action, Job, Meeting, MeetingAction, MinutesDraft, RegisterEntry, ToolSwitch, Transaction, Transcript, User
from app.schemas.minutes import MeetingActionResponse, MeetingActionReview, MeetingActionUpdate, MeetingCreate, MeetingResponse, MinutesResponse, MinutesUpdate, TranscriptOutput
from app.schemas.jobs import JobResponse
from app.services.audit import add_audit_event
from app.services.meeting_jobs import process_minutes_job
from app.services.minutes_graph import resolve_meeting
from app.services.register import approve_proposal, financial_year_for, reject_proposal

router = APIRouter(prefix="/meetings", tags=["meetings and minutes"])
READ_ROLES = ("minister", "cos", "director_td", "liaison", "legal", "analyst", "data_coordinator", "admin")
DRAFT_ROLES = ("cos", "director_td")
REVIEW_ROLES = ("cos", "director_td")


def get_ai_gateway() -> AIGateway:
    return AIGateway()


def get_stt_provider() -> STTProvider:
    return FasterWhisperProvider()


def _aware(value: datetime) -> datetime:
    return value.replace(tzinfo=timezone.utc) if value.tzinfo is None else value


def _meeting_response(db: Session, meeting_id: str) -> MeetingResponse:
    meeting = resolve_meeting(db, meeting_id)
    if meeting is None:
        raise HTTPException(status_code=404, detail="Meeting was not found")
    transcript = TranscriptOutput.model_validate(meeting.transcript) if meeting.transcript else None
    latest_minutes = meeting.minutes[0] if meeting.minutes else None
    cached_run = db.get(AIRun, latest_minutes.drafted_by_run_id) if latest_minutes and latest_minutes.drafted_by_run_id else None
    minutes_response = MinutesResponse.model_validate(latest_minutes) if latest_minutes else None
    if minutes_response is not None:
        minutes_response.demo_cache_used = cached_run is not None and cached_run.error_code == "demo_cache_used"
    actions = [MeetingActionResponse.model_validate({**{column.name: getattr(item, column.name) for column in MeetingAction.__table__.columns}, "officer_name": item.officer.name if item.officer else None}) for item in meeting.actions]
    intake_actions = []
    if meeting.source_document_id:
        intake_actions = [
            {"id": item.id, "description": item.description, "owner": item.owner_org.name if item.owner_org else item.owner_person_name,
             "deadline": item.deadline.isoformat() if item.deadline else None, "deadline_text": item.deadline_text,
             "priority": item.priority, "status": item.status, "serial_no": item.serial_no,
             "project": item.project.name if item.project else None, "source_quote": item.source_quote, "source_page": item.source_page}
            for item in db.scalars(select(Action).where(Action.source_document_id == meeting.source_document_id).order_by(Action.created_at, Action.id))
        ]
    return MeetingResponse(
        id=meeting.id, title=meeting.title, occurred_at=meeting.occurred_at, occasion_type=meeting.occasion_type,
        location=meeting.location, attendees=meeting.attendees, decisions=meeting.decisions or [], summary=meeting.summary or "",
        source_document_id=meeting.source_document_id, intake_actions=intake_actions,
        audio_url=f"/api/v1/meetings/{meeting.id}/audio" if meeting.audio_path else None,
        language_hint=meeting.language_hint, classification=meeting.classification, is_synthetic=meeting.is_synthetic,
        status=meeting.status, draft_due_at=meeting.draft_due_at, issue_due_at=meeting.issue_due_at,
        transcript=transcript, minutes=minutes_response,
        actions=actions, created_at=meeting.created_at, updated_at=meeting.updated_at,
    )


def _meeting_for_user(db: Session, meeting_id: str, actor: User) -> Meeting:
    meeting = db.get(Meeting, meeting_id)
    if meeting is None or (meeting.classification == "confidential" and actor.role not in {"minister", "cos"}):
        raise HTTPException(status_code=404, detail="Meeting was not found")
    return meeting


@router.get("", response_model=list[MeetingResponse])
def list_meetings(db: Session = Depends(get_db), actor: User = Depends(require_roles(*READ_ROLES))) -> list[MeetingResponse]:
    ids = db.scalars(select(Meeting.id).order_by(Meeting.occurred_at.desc())).all()
    rows = [db.get(Meeting, meeting_id) for meeting_id in ids]
    visible = [row.id for row in rows if row is not None and (row.classification != "confidential" or actor.role in {"minister", "cos"})]
    return [_meeting_response(db, meeting_id) for meeting_id in visible]


@router.get("/{meeting_id}", response_model=MeetingResponse)
def get_meeting(meeting_id: str, db: Session = Depends(get_db), actor: User = Depends(require_roles(*READ_ROLES))) -> MeetingResponse:
    _meeting_for_user(db, meeting_id, actor)
    return _meeting_response(db, meeting_id)


@router.post("", response_model=MeetingResponse, status_code=201)
def create_meeting(payload: MeetingCreate, request: Request, db: Session = Depends(get_db), actor: User = Depends(require_roles(*DRAFT_ROLES))) -> MeetingResponse:
    occurred_at = _aware(payload.occurred_at)
    if payload.classification == "confidential" and actor.role != "cos":
        raise HTTPException(status_code=403, detail="Only the CoS may create a confidential meeting record")
    now = datetime.now(timezone.utc)
    meeting = Meeting(title=payload.title.strip(), occurred_at=occurred_at, occasion_type=payload.occasion_type, location=payload.location.strip(), attendees=[item.model_dump() for item in payload.attendees], language_hint=payload.language_hint, classification=payload.classification, is_synthetic=False, status="uploaded", draft_due_at=occurred_at + timedelta(hours=24), issue_due_at=occurred_at + timedelta(hours=48), created_by=actor.id)
    db.add(meeting)
    db.flush()
    add_audit_event(db, actor, request, action="meeting_created", entity_type="meeting", entity_id=meeting.id, after={"title": meeting.title, "occurred_at": occurred_at.isoformat(), "classification": meeting.classification})
    db.commit()
    return _meeting_response(db, meeting.id)


@router.post("/{meeting_id}/audio", response_model=MeetingResponse)
async def upload_audio(meeting_id: str, request: Request, file: UploadFile = File(...), db: Session = Depends(get_db), actor: User = Depends(require_roles(*DRAFT_ROLES))) -> MeetingResponse:
    meeting = _meeting_for_user(db, meeting_id, actor)
    if meeting.status == "issued":
        raise HTTPException(status_code=409, detail="Issued meeting records cannot be changed")
    suffix = Path(file.filename or "").suffix.lower()
    if suffix not in {".wav", ".mp3", ".m4a", ".mp4", ".webm", ".ogg", ".flac"}:
        raise HTTPException(status_code=415, detail="Upload a WAV, MP3, M4A, MP4, WebM, OGG, or FLAC recording")
    content = await file.read(25 * 1024 * 1024 + 1)
    if not content or len(content) > 25 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="Recording must be between 1 byte and 25 MB")
    audio_dir = settings.resolved_data_dir / "audio"
    audio_dir.mkdir(parents=True, exist_ok=True)
    relative = Path("audio") / f"{uuid4()}{suffix}"
    destination = settings.resolved_data_dir / relative
    destination.write_bytes(content)
    previous = meeting.audio_path
    meeting.audio_path = str(relative)
    meeting.status = "uploaded"
    add_audit_event(db, actor, request, action="meeting_audio_uploaded", entity_type="meeting", entity_id=meeting.id, before={"audio_path": previous}, after={"audio_path": str(relative), "byte_count": len(content)})
    db.commit()
    return _meeting_response(db, meeting.id)


@router.get("/{meeting_id}/audio")
def get_audio(meeting_id: str, db: Session = Depends(get_db), actor: User = Depends(require_roles(*READ_ROLES))) -> FileResponse:
    meeting = _meeting_for_user(db, meeting_id, actor)
    if not meeting.audio_path:
        raise HTTPException(status_code=404, detail="Meeting recording was not found")
    path = (settings.resolved_data_dir / meeting.audio_path).resolve()
    if settings.resolved_data_dir not in path.parents or not path.is_file():
        raise HTTPException(status_code=404, detail="Meeting recording was not found")
    return FileResponse(path, filename=f"meeting-{meeting.id[:8]}{path.suffix}")


@router.post("/{meeting_id}/process", response_model=JobResponse, status_code=202)
def process_meeting(meeting_id: str, background_tasks: BackgroundTasks, request: Request, db: Session = Depends(get_db), actor: User = Depends(require_roles(*DRAFT_ROLES)), gateway: AIGateway = Depends(get_ai_gateway), stt: STTProvider = Depends(get_stt_provider)) -> JobResponse:
    meeting = _meeting_for_user(db, meeting_id, actor)
    if meeting.status == "issued":
        raise HTTPException(status_code=409, detail="Issued meeting records cannot be redrafted")
    if meeting.status == "draft_ready" or meeting.status == "under_review":
        raise HTTPException(status_code=409, detail="A minutes draft already exists; edit or review the current draft")
    active = db.scalar(select(Job.id).where(Job.resource_type == "meeting", Job.resource_id == meeting_id, Job.status.in_(("queued", "running"))).limit(1))
    if active:
        raise HTTPException(status_code=409, detail="Meeting processing is already queued or running")
    switch = db.get(ToolSwitch, "minutes_assistant")
    if switch is None or switch.suspended:
        gateway.record_blocked_attempt(db, tool="minutes_assistant", role="extraction", classification=meeting.classification, is_synthetic=meeting.is_synthetic, user_id=actor.id, prompt_version="minutes-v1", error_code="tool_switch_missing" if switch is None else "tool_suspended")
        raise HTTPException(status_code=503, detail="Minutes Assistant is suspended by an AI assurance control")
    job = Job(owner_id=actor.id, resource_type="meeting", resource_id=meeting_id, status="queued", progress=0, message="Waiting to process meeting")
    db.add(job)
    db.flush()
    add_audit_event(db, actor, request, action="meeting_processing_queued", entity_type="job", entity_id=job.id, after={"meeting_id": meeting_id})
    db.commit()
    db.refresh(job)
    background_tasks.add_task(process_minutes_job, db, job.id, gateway, stt)
    return JobResponse.model_validate(job)


@router.post("/{meeting_id}/review", response_model=MeetingResponse)
def start_review(meeting_id: str, request: Request, db: Session = Depends(get_db), actor: User = Depends(require_roles(*REVIEW_ROLES))) -> MeetingResponse:
    meeting = _meeting_for_user(db, meeting_id, actor)
    if meeting.status == "draft_ready":
        meeting.status = "under_review"
        add_audit_event(db, actor, request, action="minutes_review_started", entity_type="meeting", entity_id=meeting_id, before={"status": "draft_ready"}, after={"status": "under_review"})
        db.commit()
    return _meeting_response(db, meeting_id)


@router.patch("/{meeting_id}/minutes", response_model=MeetingResponse)
def update_minutes(meeting_id: str, payload: MinutesUpdate, request: Request, db: Session = Depends(get_db), actor: User = Depends(require_roles(*REVIEW_ROLES))) -> MeetingResponse:
    meeting = _meeting_for_user(db, meeting_id, actor)
    if meeting.status == "issued":
        raise HTTPException(status_code=404 if meeting is None else 409, detail="Meeting draft was not found or is already issued")
    draft = db.scalar(select(MinutesDraft).where(MinutesDraft.meeting_id == meeting_id).order_by(MinutesDraft.version.desc()).limit(1))
    if draft is None:
        raise HTTPException(status_code=409, detail="Draft minutes have not been generated")
    before = {"body_markdown": draft.body_markdown}
    draft.body_markdown = payload.body_markdown.strip()
    draft.status = "under_review"
    meeting.status = "under_review"
    run = db.get(AIRun, draft.drafted_by_run_id) if draft.drafted_by_run_id else None
    if run:
        run.diff_summary = "Minutes text edited by a named reviewer; final text recorded on issue."
    add_audit_event(db, actor, request, action="minutes_edited", entity_type="minutes", entity_id=draft.id, before=before, after={"body_markdown": draft.body_markdown})
    db.commit()
    return _meeting_response(db, meeting_id)


def _ensure_register_proposal(db: Session, meeting: Meeting, action: MeetingAction) -> RegisterEntry:
    if action.officer_id is None or action.deadline_date is None:
        raise HTTPException(status_code=422, detail="Resolve the action owner and deadline before approving it")
    entry = db.get(RegisterEntry, action.register_entry_id) if action.register_entry_id else None
    if entry is None:
        entry = RegisterEntry(fy=financial_year_for(meeting.occurred_at.date()), serial_no=None, occurred_at=meeting.occurred_at, occasion_type=meeting.occasion_type, occasion_ref=meeting.title, description=action.description, officer_id=action.officer_id, deadline=action.deadline_date, priority=action.priority_suggestion, source_record=f"MIN-{meeting.id[:8]}", status="open", transaction_id=action.transaction_id, origin="ai_proposed", review_state="proposed", ai_assistance={"tool": "minutes_assistant", "run_id": action.ai_run_id}, evidence={"quote": action.quote, "transcript_start": action.transcript_start, "transcript_end": action.transcript_end, "confidence": action.confidence})
        db.add(entry)
        db.flush()
        action.register_entry_id = entry.id
    else:
        entry.description = action.description
        entry.officer_id = action.officer_id
        entry.deadline = action.deadline_date
        entry.priority = action.priority_suggestion
        entry.transaction_id = action.transaction_id
    return entry


@router.patch("/{meeting_id}/actions/{action_id}", response_model=MeetingActionResponse)
def update_action(meeting_id: str, action_id: str, payload: MeetingActionUpdate, request: Request, db: Session = Depends(get_db), actor: User = Depends(require_roles(*REVIEW_ROLES))) -> MeetingActionResponse:
    meeting = _meeting_for_user(db, meeting_id, actor)
    action = db.get(MeetingAction, action_id)
    if action is None or action.meeting_id != meeting_id:
        raise HTTPException(status_code=404, detail="Meeting action was not found")
    if action.review_state != "pending" or meeting.status == "issued":
        raise HTTPException(status_code=409, detail="Resolved or issued actions cannot be edited")
    before = jsonable_encoder({key: getattr(action, key) for key in payload.model_fields_set})
    changes = payload.model_dump(exclude_unset=True)
    if "officer_id" in changes:
        officer = db.get(User, changes["officer_id"]) if changes["officer_id"] else None
        if officer is None or not officer.active:
            raise HTTPException(status_code=422, detail="Select an active officer")
    if "deadline_date" in changes and changes["deadline_date"] is not None and changes["deadline_date"] < meeting.occurred_at.date():
        raise HTTPException(status_code=422, detail="Action deadline cannot be earlier than the meeting")
    if "transaction_id" in changes and changes["transaction_id"] and db.get(Transaction, changes["transaction_id"]) is None:
        raise HTTPException(status_code=422, detail="Linked transaction was not found")
    for key, value in changes.items():
        setattr(action, key, value)
    if action.officer_id and action.deadline_date:
        action.review_reason = None
        entry = _ensure_register_proposal(db, meeting, action)
        add_audit_event(db, actor, request, action="minutes_action_updated", entity_type="meeting_action", entity_id=action.id, before=before, after=jsonable_encoder(changes | {"register_entry_id": entry.id}))
    else:
        missing = []
        if action.officer_id is None:
            missing.append("Confirm the responsible officer")
        if action.deadline_date is None:
            missing.append("Set a valid deadline")
        action.review_reason = "; ".join(missing)
        add_audit_event(db, actor, request, action="minutes_action_updated", entity_type="meeting_action", entity_id=action.id, before=before, after=jsonable_encoder(changes))
    db.commit()
    officer = db.get(User, action.officer_id) if action.officer_id else None
    return MeetingActionResponse.model_validate({**{column.name: getattr(action, column.name) for column in MeetingAction.__table__.columns}, "officer_name": officer.name if officer else None})


@router.post("/{meeting_id}/actions/{action_id}/review", response_model=MeetingActionResponse)
def review_action(meeting_id: str, action_id: str, payload: MeetingActionReview, request: Request, db: Session = Depends(get_db), actor: User = Depends(require_roles(*REVIEW_ROLES))) -> MeetingActionResponse:
    meeting = _meeting_for_user(db, meeting_id, actor)
    action = db.get(MeetingAction, action_id)
    if action is None or action.meeting_id != meeting_id:
        raise HTTPException(status_code=404, detail="Meeting action was not found")
    if action.review_state != "pending" or meeting.status == "issued":
        raise HTTPException(status_code=409, detail="This action has already been resolved")
    before = {"review_state": action.review_state, "review_reason": action.review_reason}
    if payload.decision == "approve":
        entry = _ensure_register_proposal(db, meeting, action)
        entry, change = approve_proposal(db, entry)
        run = db.get(AIRun, action.ai_run_id) if action.ai_run_id else None
        if run:
            run.approver_id = actor.id
            run.approved_at = datetime.now(timezone.utc)
            run.diff_summary = "Action checked and approved by a named officer."
        action.review_state = "approved"
        action.register_entry_id = entry.id
        action.review_reason = None
        add_audit_event(db, actor, request, action="minutes_action_approved", entity_type="meeting_action", entity_id=action.id, before=before, after={"review_state": action.review_state, "register_entry_id": entry.id, "serial_no": entry.serial_no})
        add_audit_event(db, actor, request, action="register_proposal_approved", entity_type="register_entry", entity_id=entry.id, before=change["before"], after=change["after"])
    else:
        if not payload.reason or len(payload.reason.strip()) < 3:
            raise HTTPException(status_code=422, detail="A rejection reason is required")
        action.review_state = "rejected"
        action.review_reason = (payload.reason or "Rejected during human review.").strip()
        entry = db.get(RegisterEntry, action.register_entry_id) if action.register_entry_id else None
        if entry is not None and entry.review_state == "proposed":
            entry, change = reject_proposal(db, entry, action.review_reason)
            add_audit_event(db, actor, request, action="register_proposal_rejected", entity_type="register_entry", entity_id=entry.id, before=change["before"], after=change["after"] | {"reason": change["reason"]})
        run = db.get(AIRun, action.ai_run_id) if action.ai_run_id else None
        if run:
            run.approver_id = actor.id
            run.approved_at = datetime.now(timezone.utc)
            run.diff_summary = f"Action rejected. Reason: {action.review_reason}"
        add_audit_event(db, actor, request, action="minutes_action_rejected", entity_type="meeting_action", entity_id=action.id, before=before, after={"review_state": action.review_state, "reason": action.review_reason})
    db.commit()
    officer = db.get(User, action.officer_id) if action.officer_id else None
    return MeetingActionResponse.model_validate({**{column.name: getattr(action, column.name) for column in MeetingAction.__table__.columns}, "officer_name": officer.name if officer else None})


@router.post("/{meeting_id}/issue", response_model=MeetingResponse)
def issue_minutes(meeting_id: str, request: Request, db: Session = Depends(get_db), actor: User = Depends(require_roles("cos"))) -> MeetingResponse:
    meeting = db.get(Meeting, meeting_id)
    if meeting is None:
        raise HTTPException(status_code=404, detail="Meeting was not found")
    if meeting.status == "issued":
        raise HTTPException(status_code=409, detail="Minutes have already been issued")
    actions = db.scalars(select(MeetingAction).where(MeetingAction.meeting_id == meeting_id)).all()
    if any(action.review_state == "pending" for action in actions):
        raise HTTPException(status_code=409, detail="Resolve every proposed action before issuing these minutes")
    draft = db.scalar(select(MinutesDraft).where(MinutesDraft.meeting_id == meeting_id).order_by(MinutesDraft.version.desc()).limit(1))
    if draft is None:
        raise HTTPException(status_code=409, detail="Draft minutes have not been generated")
    now = datetime.now(timezone.utc)
    draft.status = "issued"
    draft.reviewed_by = actor.id
    draft.issued_at = now
    meeting.status = "issued"
    run = db.get(AIRun, draft.drafted_by_run_id) if draft.drafted_by_run_id else None
    if run:
        run.approver_id = actor.id
        run.approved_at = now
        run.diff_summary = f"Minutes issued after reviewing {len(actions)} extracted action(s)."
    exports_dir = settings.resolved_data_dir / "exports"
    exports_dir.mkdir(parents=True, exist_ok=True)
    export_file = exports_dir / f"minutes-{meeting.id}.docx"
    document = Document()
    document.add_heading(meeting.title, level=0)
    document.add_paragraph(f"Date: {meeting.occurred_at:%d %B %Y} | Place: {meeting.location}")
    document.add_paragraph(f"Attendees: {', '.join(item.get('name', '') for item in meeting.attendees)}")
    document.add_paragraph(f"Status: Issued by {actor.name} on {now:%d %B %Y %H:%M UTC}")
    for line in draft.body_markdown.splitlines():
        if line.startswith("# "):
            document.add_heading(line[2:].strip(), level=1)
        elif line.startswith("## "):
            document.add_heading(line[3:].strip(), level=2)
        elif line.strip():
            document.add_paragraph(line.strip())
    document.save(export_file)
    draft.export_path = str(export_file)
    add_audit_event(db, actor, request, action="minutes_issued", entity_type="minutes", entity_id=draft.id, before={"status": "under_review"}, after={"status": "issued", "issued_at": now.isoformat(), "action_count": len(actions)})
    db.commit()
    return _meeting_response(db, meeting_id)


@router.get("/{meeting_id}/minutes.docx")
def download_minutes(meeting_id: str, db: Session = Depends(get_db), _: User = Depends(require_roles(*READ_ROLES))) -> FileResponse:
    meeting = db.get(Meeting, meeting_id)
    draft = db.scalar(select(MinutesDraft).where(MinutesDraft.meeting_id == meeting_id, MinutesDraft.status == "issued").order_by(MinutesDraft.version.desc()).limit(1))
    if meeting is None or draft is None or not draft.export_path:
        raise HTTPException(status_code=404, detail="Issued minutes export was not found")
    path = Path(draft.export_path).resolve()
    if settings.resolved_data_dir not in path.parents or not path.is_file():
        raise HTTPException(status_code=404, detail="Issued minutes export was not found")
    return FileResponse(path, filename=f"minutes-{meeting.id[:8]}.docx", media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
