from __future__ import annotations

import asyncio
import hashlib
import math
import re
from functools import lru_cache
from datetime import datetime, time, timedelta, timezone
from pathlib import Path
from uuid import uuid4

import chromadb
from docx import Document as WordDocument
from pypdf import PdfReader
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.ai.gateway import AIGateway
from app.core.config import settings
from app.db.models import Document, DocumentChunk, Meeting, MinutesDraft, RegisterEntry, Transaction, User

DIMENSIONS = 384
COLLECTION = "pc_command_document_chunks"
TOKEN_RE = re.compile(r"[\w'-]+", re.UNICODE)


def tokenize(text: str) -> list[str]:
    return [token.casefold() for token in TOKEN_RE.findall(text)]


def embed_local(text: str) -> list[float]:
    """Embed with the configured multilingual model, or a deterministic offline fallback."""
    model_name = settings.embedding_model.strip()
    if model_name:
        return embed_local_texts([text], model_name)[0]

    # Keep a no-download fallback for clean demo setups. It is deliberately distinct
    # from semantic multilingual embeddings and is identified as such in the index.
    vector = [0.0] * DIMENSIONS
    tokens = tokenize(text)
    for index, token in enumerate(tokens):
        digest = hashlib.blake2b(token.encode("utf-8"), digest_size=8).digest()
        bucket = int.from_bytes(digest[:4], "little") % DIMENSIONS
        sign = 1.0 if digest[4] & 1 else -1.0
        vector[bucket] += sign * (1.0 + math.log1p(len(tokens) / (index + 1)))
    norm = math.sqrt(sum(value * value for value in vector)) or 1.0
    return [value / norm for value in vector]


@lru_cache(maxsize=2)
def _sentence_transformer(model_name: str):
    try:
        from sentence_transformers import SentenceTransformer
    except ImportError as exc:
        raise RuntimeError(
            "Install backend requirements to use the configured sentence-transformers embedding model"
        ) from exc
    return SentenceTransformer(model_name)


def embed_local_texts(texts: list[str], model_name: str) -> list[list[float]]:
    """Return normalized multilingual sentence-transformer vectors in input order."""
    if not texts:
        return []
    vectors = _sentence_transformer(model_name).encode(
        texts,
        normalize_embeddings=True,
        convert_to_numpy=True,
        show_progress_bar=False,
    )
    return vectors.tolist()


def embed_texts(db: Session, texts: list[str], *, classification: str, is_synthetic: bool, user_id: str | None) -> list[list[float]]:
    if settings.embedding_provider == "openrouter" or settings.model_embedding.strip():
        return asyncio.run(AIGateway().embed(db, tool="records_search", texts=texts, classification=classification, is_synthetic=is_synthetic, user_id=user_id))
    if settings.embedding_provider != "local":
        raise ValueError("EMBEDDING_PROVIDER must be 'local' or 'openrouter'")
    if settings.embedding_model.strip():
        return embed_local_texts(texts, settings.embedding_model.strip())
    return [embed_local(text) for text in texts]


def chroma_collection():
    path = settings.resolved_data_dir / "chroma"
    path.mkdir(parents=True, exist_ok=True)
    client = chromadb.PersistentClient(path=str(path))
    provider_zone = settings.ai_provider_embedding.strip() or settings.ai_zone_mode
    local_model = settings.embedding_model.strip() or "local-hash-v1"
    embedding_key = f"{provider_zone}:{settings.model_embedding.strip() or local_model}"
    collection_suffix = hashlib.sha1(embedding_key.encode("utf-8")).hexdigest()[:8]
    return client.get_or_create_collection(f"{COLLECTION}_{collection_suffix}", metadata={"hnsw:space": "cosine", "embedding_key": embedding_key})


def extract_pages(path: Path) -> list[tuple[int | None, str]]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        reader = PdfReader(str(path))
        return [(index + 1, page.extract_text() or "") for index, page in enumerate(reader.pages)]
    if suffix == ".docx":
        doc = WordDocument(str(path))
        return [(None, "\n".join(p.text for p in doc.paragraphs if p.text.strip()))]
    if suffix in {".txt", ".md"}:
        return [(None, path.read_text(encoding="utf-8-sig"))]
    raise ValueError("Supported formats are PDF, DOCX, TXT, and MD")


def section_for(text: str) -> str | None:
    headings = [line.strip(" #\t") for line in text.splitlines() if line.strip().startswith("#")]
    return headings[-1][:255] if headings else None


def split_chunks(pages: list[tuple[int | None, str]], chunk_chars: int = 3200, overlap: int = 320) -> list[dict]:
    chunks: list[dict] = []
    for page, text in pages:
        normalized = re.sub(r"\n{3,}", "\n\n", text).strip()
        start = 0
        while start < len(normalized):
            end = min(len(normalized), start + chunk_chars)
            if end < len(normalized):
                boundary = normalized.rfind("\n", start + chunk_chars // 2, end)
                if boundary > start:
                    end = boundary
            body = normalized[start:end].strip()
            if body:
                chunks.append({"text": body, "page": page, "section": section_for(body), "char_start": start, "char_end": end})
            if end >= len(normalized):
                break
            start = max(start + 1, end - overlap)
    return chunks


def index_structured_records(db: Session) -> None:
    """Materialize approved Register entries and issued minutes as searchable sources."""
    records: list[dict] = []
    entries = db.scalars(select(RegisterEntry).where(RegisterEntry.review_state == "approved")).all()
    for entry in entries:
        note = f"structured:register:{entry.id}"
        if db.scalar(select(Document.id).where(Document.source_note == note)):
            continue
        officer = db.get(User, entry.officer_id)
        transaction = db.get(Transaction, entry.transaction_id) if entry.transaction_id else None
        label = f"Register {entry.fy}/{entry.serial_no:03d}" if entry.serial_no else "Register action"
        text = f"{label}\nDate: {entry.occurred_at:%d %B %Y}\nAction: {entry.description}\nResponsible officer: {officer.name if officer else 'Unassigned'}\nDeadline: {entry.deadline:%d %B %Y}\nPriority: {entry.priority}\nStatus: {entry.status}\nTransaction: {transaction.code if transaction else 'None'}"
        records.append({"note": note, "title": label, "doc_type": "register", "classification": "internal", "transaction_id": entry.transaction_id, "synthetic": False, "text": text})
    issued = db.scalars(select(MinutesDraft).where(MinutesDraft.status == "issued")).all()
    for draft in issued:
        note = f"structured:minutes:{draft.id}"
        if db.scalar(select(Document.id).where(Document.source_note == note)):
            continue
        meeting = db.get(Meeting, draft.meeting_id)
        if meeting is None:
            continue
        text = f"{meeting.title}\nDate: {meeting.occurred_at:%d %B %Y}\nPlace: {meeting.location}\nAttendees: {', '.join(item.get('name', '') for item in meeting.attendees)}\n\n{draft.body_markdown}"
        records.append({"note": note, "title": f"Issued minutes: {meeting.title}", "doc_type": "minute", "classification": meeting.classification, "transaction_id": None, "synthetic": meeting.is_synthetic, "text": text})
    if not records:
        return
    storage = settings.resolved_data_dir / "documents" / "structured"
    storage.mkdir(parents=True, exist_ok=True)
    all_ids: list[str] = []
    embedding_groups: dict[tuple[str, bool], list[int]] = {}
    all_texts: list[str] = []
    all_metadatas: list[dict] = []
    collection = chroma_collection()
    for item in records:
        saved_path = storage / f"{item['note'].replace(':', '-')}.md"
        saved_path.write_text(item["text"], encoding="utf-8")
        document = Document(title=item["title"], doc_type=item["doc_type"], classification=item["classification"], transaction_id=item["transaction_id"], file_path=str(saved_path), text_extracted=item["text"], source_note=item["note"], is_synthetic=item["synthetic"], ingest_status="ready")
        db.add(document)
        db.flush()
        for index, chunk in enumerate(split_chunks([(None, item["text"])])):
            vector_id = str(uuid4())
            db.add(DocumentChunk(document_id=document.id, chunk_index=index, vector_id=vector_id, **chunk))
            all_ids.append(vector_id)
            all_texts.append(chunk["text"])
            all_metadatas.append({"document_id": document.id, "classification": item["classification"], "doc_type": item["doc_type"], "transaction_id": item["transaction_id"] or ""})
            embedding_groups.setdefault((item["classification"], item["synthetic"]), []).append(len(all_texts) - 1)
    all_vectors: list[list[float] | None] = [None] * len(all_texts)
    for (classification, synthetic), indexes in embedding_groups.items():
        vectors = embed_texts(db, [all_texts[index] for index in indexes], classification=classification, is_synthetic=synthetic, user_id=None)
        for index, vector in zip(indexes, vectors):
            all_vectors[index] = vector
    collection.add(ids=all_ids, embeddings=[vector for vector in all_vectors if vector is not None], documents=all_texts, metadatas=all_metadatas)
    db.commit()


def ingest_document(db: Session, *, title: str, doc_type: str, classification: str, transaction_id: str | None, source_note: str | None, is_synthetic: bool, saved_path: Path, user_id: str | None = None) -> Document:
    pages = extract_pages(saved_path)
    text = "\n\n".join(body for _, body in pages).strip()
    if not text:
        raise ValueError("No text could be extracted from this document")
    document = Document(title=title, doc_type=doc_type, classification=classification, transaction_id=transaction_id, file_path=str(saved_path), text_extracted=text, source_note=source_note, is_synthetic=is_synthetic, ingest_status="processing")
    db.add(document)
    db.flush()
    rows = split_chunks(pages)
    if not rows:
        raise ValueError("No searchable text chunks were created")
    collection = chroma_collection()
    ids: list[str] = []
    metadatas: list[dict] = []
    for index, row in enumerate(rows):
        vector_id = str(uuid4())
        db.add(DocumentChunk(document_id=document.id, chunk_index=index, vector_id=vector_id, **row))
        ids.append(vector_id)
        metadatas.append({"document_id": document.id, "classification": classification, "doc_type": doc_type, "transaction_id": transaction_id or ""})
    vectors = embed_texts(db, [row["text"] for row in rows], classification=classification, is_synthetic=is_synthetic, user_id=user_id)
    collection.add(ids=ids, embeddings=vectors, documents=[row["text"] for row in rows], metadatas=metadatas)
    document.ingest_status = "ready"
    db.commit()
    db.refresh(document)
    return document


def index_intake_document(db: Session, document: Document) -> None:
    """Index an already saved intake source using local embeddings only."""
    if db.scalar(select(DocumentChunk.id).where(DocumentChunk.document_id == document.id)):
        return
    rows = split_chunks([(1, document.text_extracted)])
    if not rows:
        raise ValueError("No searchable text chunks were created")
    collection = chroma_collection()
    ids = [str(uuid4()) for _ in rows]
    vectors = [embed_local(row["text"]) for row in rows]
    db.add_all([DocumentChunk(document_id=document.id, chunk_index=index, vector_id=ids[index], **row) for index, row in enumerate(rows)])
    collection.add(
        ids=ids,
        embeddings=vectors,
        documents=[row["text"] for row in rows],
        metadatas=[{"document_id": document.id, "classification": document.classification, "doc_type": document.doc_type, "transaction_id": ""} for _ in rows],
    )


def lexical_scores(query: str, chunks: list[DocumentChunk]) -> dict[str, float]:
    terms = tokenize(query)
    docs = [tokenize(chunk.text) for chunk in chunks]
    n = len(docs)
    avgdl = sum(map(len, docs)) / max(n, 1)
    scores: dict[str, float] = {}
    for chunk, tokens in zip(chunks, docs):
        counts = {term: tokens.count(term) for term in set(terms)}
        score = 0.0
        for term, tf in counts.items():
            df = sum(term in doc for doc in docs)
            idf = math.log(1 + (n - df + 0.5) / (df + 0.5))
            score += idf * (tf * 2.2) / (tf + 1.2 * (0.25 + 0.75 * len(tokens) / max(avgdl, 1)))
        scores[chunk.vector_id] = score
    return scores


def retrieve(db: Session, *, query: str, classification_access: list[str], transaction_id: str | None = None, doc_type: str | None = None, date_from=None, date_to=None, top_k: int = 5, user_id: str | None = None) -> list[tuple[DocumentChunk, float]]:
    statement = select(DocumentChunk).join(Document).where(Document.ingest_status == "ready", Document.classification.in_(classification_access))
    if transaction_id:
        statement = statement.where(Document.transaction_id == transaction_id)
    if doc_type:
        statement = statement.where(Document.doc_type == doc_type)
    if date_from:
        statement = statement.where(Document.created_at >= datetime.combine(date_from, time.min, tzinfo=timezone.utc))
    if date_to:
        statement = statement.where(Document.created_at < datetime.combine(date_to + timedelta(days=1), time.min, tzinfo=timezone.utc))
    chunks = db.scalars(statement).all()
    if not chunks:
        return []
    allowed_ids = [chunk.vector_id for chunk in chunks]
    collection = chroma_collection()
    vector_ids = collection.get(ids=allowed_ids, include=[]).get("ids", [])
    vector_rank: dict[str, int] = {}
    if vector_ids:
        classification_order = {"public": 0, "internal": 1, "confidential": 2}
        max_classification = max((chunk.document.classification for chunk in chunks), key=lambda item: classification_order[item])
        synthetic = all(chunk.document.is_synthetic for chunk in chunks)
        query_vector = embed_texts(db, [query], classification=max_classification, is_synthetic=synthetic, user_id=user_id)[0]
        result = collection.query(query_embeddings=[query_vector], ids=vector_ids, n_results=min(len(vector_ids), max(top_k * 4, 20)), include=["distances"])
        vector_rank = {identifier: rank for rank, identifier in enumerate((result.get("ids") or [[]])[0], start=1)}
    lexical = lexical_scores(query, chunks)
    lexical_rank = {identifier: rank for rank, (identifier, _) in enumerate(sorted(lexical.items(), key=lambda item: item[1], reverse=True), start=1)}
    by_id = {chunk.vector_id: chunk for chunk in chunks}
    fused = {identifier: (1 / (60 + vector_rank[identifier]) if identifier in vector_rank else 0) + (1 / (60 + lexical_rank[identifier]) if lexical.get(identifier, 0) > 0 else 0) for identifier in allowed_ids}
    return [(by_id[identifier], score) for identifier, score in sorted(fused.items(), key=lambda item: item[1], reverse=True)[:top_k] if score > 0]
